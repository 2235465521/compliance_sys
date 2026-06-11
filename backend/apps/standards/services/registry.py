from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
import unicodedata

from django.core.files.uploadedfile import UploadedFile
from django.db import IntegrityError
from django.db.models import Count, Q
from django.db.models.functions import Trim
from django.db.models.query import QuerySet
from django.http import HttpRequest

from apps.standards.models import NationalStandardBasic, NationalStandardExtension
from apps.standards.schemas.registry import (
    ApiEnvelopeStatistics,
    BatchImportOut,
    StandardDetailOut,
    StandardExtensionOut,
    StandardListItemOut,
    StandardPaginatedOut,
    StandardPatchIn,
    StandardStatisticsData,
    StatusCountItem,
    StatusGroupItem,
    YearCountItem,
)


class DuplicateStdCodeError(ValueError):
    """on_duplicate=reject 且库中已有该 std_code。"""

    def __init__(self, code: str):
        self.code = code
        super().__init__(f"std_code 已存在，未写入（避免覆盖）: {code}")


def _normalize_std_code_import(s: str) -> str:
    """Excel/CSV 国标号：NFC、去 BOM/零宽/NBSP，折叠空白。"""
    if not s:
        return ""
    t = unicodedata.normalize("NFC", str(s).strip())
    for ch in ("\ufeff", "\u200b", "\u200c", "\u200d", "\xa0"):
        t = t.replace(ch, "")
    return " ".join(t.split())


def _find_basic_by_std_code_variants(code: str) -> NationalStandardBasic | None:
    """精确匹配 + TRIM(std_code) 匹配，避免库字段尾部空格等导致『已存在』却查不到。"""
    bid = _normalize_std_code_import(code)
    if not bid:
        return None
    normalized = " ".join(bid.split())
    compact = normalized.replace(" ", "")
    candidates: list[str] = []
    for c in (bid, normalized, compact):
        if c and c not in candidates:
            candidates.append(c)

    for c in candidates:
        o = NationalStandardBasic.objects.filter(std_code=c).first()
        if o:
            return o

    trimmed = (
        NationalStandardBasic.objects.annotate(_trim_code=Trim("std_code"))
        .filter(_trim_code__in=candidates)
        .first()
    )
    if trimmed:
        return trimmed
    return None


_HEADER_ALIASES: dict[str, str] = {
    "国标号": "std_code",
    "标准号": "std_code",
    "std_code": "std_code",
    "标准名称": "std_name",
    "std_name": "std_name",
    "标准状态": "std_status",
    "标准状态（Excel 原文，如废止、现行）": "std_status",
    "std_status": "std_status",
    "publish_date": "publish_date",
    "发布日期": "publish_date",
    "effective_date": "effective_date",
    "实施日期": "effective_date",
    "abolition_date": "abolition_date",
    "废止日期": "abolition_date",
    "标准类别": "std_category",
    "std_category": "std_category",
    "代替标准": "replaces_std_code",
    "replaces_std_code": "replaces_std_code",
    "代替类型": "replace_type",
    "代替类型：-1无；1全部代替；2部分代替；3部分代完；4未知": "replace_type",
    "replace_type": "replace_type",
    "中国标准分类号": "ccs_code",
    "ccs_code": "ccs_code",
    "国际标准分类号": "ics_code",
    "ics_code": "ics_code",
    "谱系号": "ped_id",
    "谱系号（可能多条，英文逗号拼接）": "ped_id",
    "ped_id": "ped_id",
    "详情链接": "detail_url",
    "detail_url": "detail_url",
    "文件路径": "std_file_path",
    "std_file_path": "std_file_path",
    "国标文件保存路径": "std_file_path",
    "ex_state": "std_status",
    "扩展状态": "std_status",
    "入库状态": "std_status",
}


def _normalize_header_key(raw: str | None) -> str:
    """表头可能带 \\ufeff 或两侧空白（Excel/UTF-8）。"""
    if not raw:
        return ""
    return (raw.replace("\ufeff", "")).strip()


def _normalize_header(cell: str) -> str | None:
    key = _normalize_header_key(cell)
    if not key:
        return None
    return _HEADER_ALIASES.get(key) or _HEADER_ALIASES.get(key.lower())


def _upload_cell_to_string(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    s = str(v).strip()
    return s if s else None


def _iter_xlsx_rows(raw: bytes):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header_row = next(rows, None)
        if not header_row:
            return
        col_map: dict[int, str] = {}
        for idx, cell in enumerate(header_row):
            if cell is None:
                continue
            fn = _normalize_header_key(str(cell))
            norm = _normalize_header(fn)
            if norm:
                col_map[idx] = norm
        for line_no, row in enumerate(rows, start=2):
            if row is None:
                continue
            if not any(v is not None and _upload_cell_to_string(v) for v in row):
                continue
            row_dict: dict[str, str] = {}
            for idx, norm in col_map.items():
                if idx < len(row):
                    s = _upload_cell_to_string(row[idx])
                    if s:
                        row_dict[norm] = s
            yield line_no, row_dict
    finally:
        wb.close()


def _parse_date(val: str | None) -> date | None:
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", s)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    return None


def _qs_filtered(search: str | None, status_filter: str | None) -> QuerySet[NationalStandardBasic]:
    qs: QuerySet[NationalStandardBasic] = NationalStandardBasic.objects.all().order_by("-id")
    if search:
        q = Q(std_code__icontains=search) | Q(std_name__icontains=search)
        qs = qs.filter(q)
    if status_filter:
        qs = qs.filter(std_status__icontains=status_filter)
    return qs


def _page_url(request: HttpRequest, page: int, page_size: int) -> str:
    q = request.GET.copy()
    q["page"] = str(page)
    q["page_size"] = str(page_size)
    return request.build_absolute_uri(f"{request.path}?{q.urlencode()}")


def list_standards(
    request: HttpRequest,
    *,
    page: int,
    page_size: int,
    search: str | None,
    status_filter: str | None,
) -> StandardPaginatedOut:
    if page < 1:
        page = 1
    if page_size < 1:
        page_size = 10
    page_size = min(page_size, 100)

    qs = _qs_filtered(search, status_filter)
    total = qs.count()
    start = (page - 1) * page_size
    rows = qs[start : start + page_size]

    next_url = _page_url(request, page + 1, page_size) if start + page_size < total else None
    prev_url = _page_url(request, page - 1, page_size) if page > 1 else None

    results = [
        StandardListItemOut(
            id=o.id,
            bz_id=o.std_code,
            bz_name=o.std_name,
            bz_release_date=o.publish_date,
            implement_time=o.effective_date,
            ex_state=o.std_status,
            replace_bz_id=o.replaces_std_code or None,
            tree_id=None,
            std_category=o.std_category,
            ccs_code=o.ccs_code,
            ics_code=o.ics_code,
            abolition_date=o.abolition_date,
            replace_type=o.replace_type,
        )
        for o in rows
    ]
    return StandardPaginatedOut(count=total, next=next_url, previous=prev_url, results=results)


def get_detail(*, bz_id: str) -> StandardDetailOut | None:
    obj: NationalStandardBasic | None = None
    if bz_id.isdigit():
        obj = NationalStandardBasic.objects.filter(id=int(bz_id)).first()
    if obj is None:
        obj = NationalStandardBasic.objects.filter(std_code=bz_id).first()
    if obj is None:
        return None

    ext: NationalStandardExtension | None = None
    try:
        ext = obj.extension_row
    except NationalStandardExtension.DoesNotExist:
        ext = None

    extension_out = None
    if ext:
        extension_out = StandardExtensionOut(
            responsible_unit=ext.responsible_unit,
            secondary_responsible_unit=ext.secondary_responsible_unit,
            issuing_department=ext.issuing_department,
            executing_unit=ext.executing_unit,
            technical_committee=ext.technical_committee,
            governing_department=ext.governing_department,
            adoption_status=ext.adoption_status,
            drafting_unit=ext.drafting_unit,
            drafter=ext.drafter,
        )

    return StandardDetailOut(
        id=obj.id,
        bz_id=obj.std_code,
        bz_name=obj.std_name,
        ex_state=obj.std_status,
        std_code=obj.std_code,
        std_name=obj.std_name,
        std_status=obj.std_status,
        publish_date=obj.publish_date,
        effective_date=obj.effective_date,
        abolition_date=obj.abolition_date,
        std_category=obj.std_category,
        replaces_std_code=obj.replaces_std_code,
        replace_type=obj.replace_type,
        ccs_code=obj.ccs_code,
        ics_code=obj.ics_code,
        ped_id=obj.ped_id,
        detail_url=obj.detail_url,
        std_file_path=obj.std_file_path,
        extension=extension_out,
    )


_BASIC_PATCH_ATTRS = frozenset(
    {
        "std_name",
        "std_status",
        "publish_date",
        "effective_date",
        "abolition_date",
        "std_category",
        "replaces_std_code",
        "replace_type",
        "ccs_code",
        "ics_code",
        "ped_id",
        "detail_url",
        "std_file_path",
    }
)

_EXT_PATCH_ATTRS = frozenset(
    {
        "responsible_unit",
        "secondary_responsible_unit",
        "issuing_department",
        "executing_unit",
        "technical_committee",
        "governing_department",
        "adoption_status",
        "drafting_unit",
        "drafter",
    }
)


def _assign_patch_value(obj, attr: str, value) -> None:
    if value is None:
        setattr(obj, attr, None)
        return
    if isinstance(value, str):
        s = value.strip()
        setattr(obj, attr, s if s else None)
        return
    setattr(obj, attr, value)


def patch_standard(*, bz_id: str, patch: StandardPatchIn) -> StandardDetailOut:
    """
    局部更新标准主表与扩展表；不改变 `std_code`（谱系节点与边以国标号为键）。
    """
    raw = patch.model_dump(exclude_unset=True)
    ext_in = raw.pop("extension", None)
    if "ex_state" in raw:
        if raw.get("std_status") is None:
            raw["std_status"] = raw.pop("ex_state")
        else:
            raw.pop("ex_state")
    ext_patch: dict[str, object] = {}
    if ext_in:
        ep = ext_in if isinstance(ext_in, dict) else ext_in.model_dump(exclude_unset=True)
        ext_patch = {k: v for k, v in ep.items() if k in _EXT_PATCH_ATTRS}

    if not raw and not ext_patch:
        raise ValueError("未提供任何可更新字段")

    s = (bz_id or "").strip()
    obj: NationalStandardBasic | None = None
    if s.isdigit():
        obj = NationalStandardBasic.objects.filter(id=int(s)).first()
    if obj is None:
        obj = NationalStandardBasic.objects.filter(std_code=s).first()
    if obj is None:
        raise ValueError(f"未找到标准: {bz_id!r}")

    for k, v in raw.items():
        if k not in _BASIC_PATCH_ATTRS:
            continue
        _assign_patch_value(obj, k, v)
    obj.save()

    if ext_patch:
        ext_obj, _ = NationalStandardExtension.objects.get_or_create(
            national_standard=obj,
            defaults={},
        )
        for k, v in ext_patch.items():
            _assign_patch_value(ext_obj, k, v)
        ext_obj.save()

    out = get_detail(bz_id=obj.std_code)
    if out is None:
        raise ValueError("更新后无法读取标准详情")
    return out


def _status_group_summary(
    *, total: int, qs: QuerySet[NationalStandardBasic]
) -> list[StatusGroupItem]:
    """按 std_status 文案归类：废止 / 现行 / 即将实施 / 其余为其它；互斥计数，合计等于 total。"""
    if total <= 0:
        return [
            StatusGroupItem(key="current", label="现行", count=0, ratio=0.0),
            StatusGroupItem(key="abolished", label="废止", count=0, ratio=0.0),
            StatusGroupItem(key="upcoming", label="即将实施", count=0, ratio=0.0),
            StatusGroupItem(key="other", label="其它", count=0, ratio=0.0),
        ]

    def ratio(c: int) -> float:
        return round((c / total) * 100, 2)

    abolished = qs.filter(std_status__icontains="废止").count()
    rest_non_abolished = qs.exclude(std_status__icontains="废止")
    upcoming = rest_non_abolished.filter(std_status__icontains="即将").count()
    rest2 = rest_non_abolished.exclude(std_status__icontains="即将")
    current = rest2.filter(std_status__icontains="现行").count()
    other = total - abolished - upcoming - current
    if other < 0:
        other = 0

    return [
        StatusGroupItem(key="current", label="现行", count=current, ratio=ratio(current)),
        StatusGroupItem(key="abolished", label="废止", count=abolished, ratio=ratio(abolished)),
        StatusGroupItem(key="upcoming", label="即将实施", count=upcoming, ratio=ratio(upcoming)),
        StatusGroupItem(key="other", label="其它", count=other, ratio=ratio(other)),
    ]


def statistics() -> ApiEnvelopeStatistics:
    from apps.standards.services.standard_status_sync import sync_national_standard_status

    sync_national_standard_status()
    base = NationalStandardBasic.objects.all()
    total = base.count()
    status_groups = _status_group_summary(total=total, qs=base)
    states_map = {g.label: g.count for g in status_groups}
    type_rows = (
        base.exclude(std_category__isnull=True)
        .exclude(std_category="")
        .values("std_category")
        .annotate(c=Count("id"))
        .order_by("-c")[:50]
    )
    types_map = {str(r["std_category"]): r["c"] for r in type_rows}
    by_status_rows = (
        base.values("std_status").annotate(c=Count("id")).order_by("-c")[:50]
    )
    by_status = [
        StatusCountItem(status=r["std_status"] or "（空）", count=r["c"]) for r in by_status_rows
    ]
    by_year_rows = (
        base.exclude(publish_date__isnull=True)
        .values("publish_date__year")
        .annotate(c=Count("id"))
        .order_by("publish_date__year")
    )
    by_year = [YearCountItem(year=int(r["publish_date__year"]), count=r["c"]) for r in by_year_rows]
    return ApiEnvelopeStatistics(
        code=200,
        data=StandardStatisticsData(
            total=total,
            types=types_map,
            states=states_map,
            status_groups=status_groups,
            by_status=by_status,
            by_publish_year=by_year,
        ),
    )


def _row_to_basic_fields(row: dict[str, str]) -> dict:
    out: dict = {}
    if "std_code" in row and row["std_code"]:
        nc = _normalize_std_code_import(str(row["std_code"]))
        if nc:
            out["std_code"] = nc
    for k in (
        "std_name",
        "std_status",
        "std_category",
        "replaces_std_code",
        "replace_type",
        "ccs_code",
        "ics_code",
        "ped_id",
        "detail_url",
        "std_file_path",
    ):
        if k in row and row[k] is not None and str(row[k]).strip():
            out[k] = str(row[k]).strip()
    for dk, sk in (
        ("publish_date", "publish_date"),
        ("effective_date", "effective_date"),
        ("abolition_date", "abolition_date"),
    ):
        if sk in row and row[sk]:
            d = _parse_date(str(row[sk]))
            if d:
                out[dk] = d
    return out


def _upsert_basic(
    fields: dict,
    *,
    extension: StandardExtensionOut | None,
    on_duplicate: str = "update",
) -> tuple[NationalStandardBasic, bool]:
    code = fields.get("std_code")
    if not code:
        raise ValueError("std_code 必填")

    obj = _find_basic_by_std_code_variants(code)
    if obj is not None and on_duplicate == "reject":
        raise DuplicateStdCodeError(code)
    created = obj is None
    if created:
        obj = NationalStandardBasic(std_code=code)
    for k, v in fields.items():
        if k == "std_code":
            continue
        setattr(obj, k, v)
    try:
        obj.save()
    except IntegrityError as exc:
        msg = str(exc).lower()
        if on_duplicate == "reject" and (
            "std_code" in msg or "duplicate" in msg or "unique" in msg
        ):
            raise DuplicateStdCodeError(code) from exc
        raise

    if extension and extension.model_dump(exclude_none=True):
        NationalStandardExtension.objects.update_or_create(
            national_standard=obj,
            defaults={
                "responsible_unit": extension.responsible_unit,
                "secondary_responsible_unit": extension.secondary_responsible_unit,
                "issuing_department": extension.issuing_department,
                "executing_unit": extension.executing_unit,
                "technical_committee": extension.technical_committee,
                "governing_department": extension.governing_department,
                "adoption_status": extension.adoption_status,
                "drafting_unit": extension.drafting_unit,
                "drafter": extension.drafter,
            },
        )
    return obj, created


def batch_import_metadata(
    files: list[UploadedFile],
    *,
    on_duplicate: str = "reject",
) -> BatchImportOut:
    imported = 0
    updated = 0
    rejected_duplicates = 0
    errors: list[str] = []

    for uf in files:
        name = (uf.name or "").lower()
        raw = uf.read()

        if name.endswith(".xlsx"):
            try:
                rows = _iter_xlsx_rows(raw)
            except ImportError as e:
                errors.append(f"{uf.name}: {e}")
                continue
            except Exception as e:  # noqa: BLE001
                errors.append(f"{uf.name}: 解析 xlsx 失败: {e}")
                continue
            for i, row in rows:
                try:
                    fields = _row_to_basic_fields(row)
                    if "std_code" not in fields:
                        errors.append(f"{uf.name} 第{i}行: 缺少 std_code/国标号")
                        continue
                    _, created = _upsert_basic(
                        fields, extension=None, on_duplicate=on_duplicate
                    )
                    if created:
                        imported += 1
                    else:
                        updated += 1
                except DuplicateStdCodeError as e:
                    rejected_duplicates += 1
                    errors.append(f"{uf.name} 第{i}行: {e}")
                except Exception as e:  # noqa: BLE001
                    errors.append(f"{uf.name} 第{i}行: {e}")
            continue

        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = raw.decode("gbk")
            except UnicodeDecodeError:
                errors.append(f"{uf.name}: 无法解码为 UTF-8 或 GBK")
                continue

        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
            reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        except csv.Error:
            reader = csv.DictReader(io.StringIO(text))

        fieldnames = reader.fieldnames or []
        col_map: dict[str, str] = {}
        for fn in fieldnames:
            fk = _normalize_header_key(fn)
            norm = _normalize_header(fk)
            if norm:
                col_map[fk] = norm

        for i, raw_row in enumerate(reader, start=2):
            row: dict[str, str] = {}
            mapped_nonempty = False
            for k, v in raw_row.items():
                kk = _normalize_header_key(k)
                nk = col_map.get(kk)
                if nk and v is not None:
                    s = str(v).strip()
                    if s:
                        mapped_nonempty = True
                    row[nk] = s
            if not mapped_nonempty:
                continue
            try:
                fields = _row_to_basic_fields(row)
                if "std_code" not in fields:
                    errors.append(f"{uf.name} 第{i}行: 缺少 std_code/国标号")
                    continue
                _, created = _upsert_basic(
                    fields, extension=None, on_duplicate=on_duplicate
                )
                if created:
                    imported += 1
                else:
                    updated += 1
            except DuplicateStdCodeError as e:
                rejected_duplicates += 1
                errors.append(f"{uf.name} 第{i}行: {e}")
            except Exception as e:  # noqa: BLE001
                errors.append(f"{uf.name} 第{i}行: {e}")

    return BatchImportOut(
        imported=imported,
        updated=updated,
        rejected_duplicates=rejected_duplicates,
        errors=errors[:200],
    )
