"""从 `resources/ICS(第7版）+CCS.xlsx` 解析 ICS/CCS，供管理命令入库。"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

RESOURCES_DIR = Path(__file__).resolve().parent.parent / "resources"
DEFAULT_ICS_CCS_XLSX = RESOURCES_DIR / "ICS(第7版）+CCS.xlsx"


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _int_level(v: Any) -> int | None:
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v)
    try:
        return int(str(v).strip())
    except ValueError:
        return None


def _build_ics_note(row: tuple[Any, ...], col_idx: dict[str, int]) -> str | None:
    parts: list[str] = []
    for key in ("扩充类目", "注释", "扩充与增补注释"):
        i = col_idx.get(key)
        if i is None or i >= len(row):
            continue
        s = _cell_str(row[i])
        if s:
            parts.append(s)
    return "\n".join(parts) if parts else None


def _header_map(header_row: tuple[Any, ...]) -> dict[str, int]:
    return {str(h).strip(): i for i, h in enumerate(header_row) if h is not None and str(h).strip()}


def load_ics_merged(xlsx_path: Path) -> list[dict[str, Any]]:
    """
    合并 ICS 与 Sheet2：主键为「分类号 u」；同名后者覆盖前者（增补表优先）。
    """
    merged: dict[str, dict[str, Any]] = {}
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        for sheet_name in ("ICS", "Sheet2"):
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)
            first = next(rows_iter, None)
            if not first:
                continue
            col_idx = _header_map(first)
            if "分类号u" not in col_idx:
                continue
            ic_u = col_idx["分类号u"]
            ic_name_i = col_idx.get("名称")
            ic_level_i = col_idx.get("层级")
            for row in rows_iter:
                if not row or ic_u >= len(row):
                    continue
                code = _cell_str(row[ic_u])
                if not code:
                    continue
                code = code[:64]
                name = ""
                if ic_name_i is not None and ic_name_i < len(row):
                    name = _cell_str(row[ic_name_i])
                level = None
                if ic_level_i is not None and ic_level_i < len(row):
                    level = _int_level(row[ic_level_i])
                note = _build_ics_note(row, col_idx)
                merged[code] = {
                    "ics_code": code,
                    "ics_level": level,
                    "ics_name": name or None,
                    "ics_note": note,
                }
    finally:
        wb.close()
    return list(merged.values())


def load_ccs_rows(xlsx_path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        if "CCS" not in wb.sheetnames:
            return []
        ws = wb["CCS"]
        rows_iter = ws.iter_rows(values_only=True)
        first = next(rows_iter, None)
        if not first:
            return []
        col_idx = _header_map(first)
        if "代码" not in col_idx or "名称" not in col_idx:
            return []
        ci = col_idx["代码"]
        ni = col_idx["名称"]
        pi = col_idx.get("父代码")
        ri = col_idx.get("备注")
        out: list[dict[str, Any]] = []
        for row in rows_iter:
            if not row or ci >= len(row):
                continue
            code = _cell_str(row[ci])[:64]
            if not code:
                continue
            name = _cell_str(row[ni]) if ni < len(row) else ""
            parent = None
            if pi is not None and pi < len(row):
                p = _cell_str(row[pi])
                parent = (p[:64] if p else None) or None
            note = None
            if ri is not None and ri < len(row):
                r = _cell_str(row[ri])
                note = r or None
            out.append(
                {
                    "ccs_code": code,
                    "ccs_name": name or None,
                    "parent_code": parent,
                    "ccs_note": note,
                }
            )
        return out
    finally:
        wb.close()
