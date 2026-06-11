"""与 `ics_industry_classification` / `ccs_industry_classification` 表字段对齐的模板与解析。"""
from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

from django.core.files.uploadedfile import UploadedFile

from apps.standards.services.classification_import import _cell_str, _int_level

RESOURCES_DIR = Path(__file__).resolve().parent.parent / "resources"
ICS_TABLE_TEMPLATE_CSV = RESOURCES_DIR / "ics_industry_classification_import_template.csv"
CCS_TABLE_TEMPLATE_CSV = RESOURCES_DIR / "ccs_industry_classification_import_template.csv"

# 表头与模型 verbose_name 一致，并兼容英文字段名导入
ICS_TEMPLATE_HEADERS: tuple[str, ...] = (
    "国际标准分类号",
    "层级",
    "名称",
    "注释/扩充说明",
)

CCS_TEMPLATE_HEADERS: tuple[str, ...] = (
    "中国标准分类号",
    "名称",
    "父代码",
    "备注",
)

ICS_HEADER_TO_FIELD: dict[str, str] = {
    "国际标准分类号": "ics_code",
    "ics_code": "ics_code",
    "层级": "ics_level",
    "ics_level": "ics_level",
    "名称": "ics_name",
    "ics_name": "ics_name",
    "注释/扩充说明": "ics_note",
    "ics_note": "ics_note",
}

CCS_HEADER_TO_FIELD: dict[str, str] = {
    "中国标准分类号": "ccs_code",
    "ccs_code": "ccs_code",
    "名称": "ccs_name",
    "ccs_name": "ccs_name",
    "父代码": "parent_code",
    "parent_code": "parent_code",
    "备注": "ccs_note",
    "ccs_note": "ccs_note",
}


def _field_col_map(header_row: tuple[Any, ...], header_map: dict[str, str]) -> dict[str, int]:
    out: dict[str, int] = {}
    for i, h in enumerate(header_row):
        if h is None:
            continue
        key = str(h).strip().replace("\ufeff", "")
        if not key:
            continue
        fn = header_map.get(key) or header_map.get(key.lower())
        if fn:
            out[fn] = i
    return out


def _ics_row_from_values(row: tuple[Any, ...], col: dict[str, int]) -> dict[str, Any] | None:
    if "ics_code" not in col:
        return None
    i = col["ics_code"]
    if i >= len(row):
        return None
    code = _cell_str(row[i])[:64]
    if not code:
        return None
    name = None
    if "ics_name" in col and col["ics_name"] < len(row):
        n = _cell_str(row[col["ics_name"]])
        name = n or None
    level = None
    if "ics_level" in col and col["ics_level"] < len(row):
        level = _int_level(row[col["ics_level"]])
    note = None
    if "ics_note" in col and col["ics_note"] < len(row):
        t = _cell_str(row[col["ics_note"]])
        note = t or None
    return {
        "ics_code": code,
        "ics_level": level,
        "ics_name": name,
        "ics_note": note,
    }


def _ccs_row_from_values(row: tuple[Any, ...], col: dict[str, int]) -> dict[str, Any] | None:
    if "ccs_code" not in col:
        return None
    ci = col["ccs_code"]
    if ci >= len(row):
        return None
    code = _cell_str(row[ci])[:64]
    if not code:
        return None
    name = ""
    if "ccs_name" in col and col["ccs_name"] < len(row):
        name = _cell_str(row[col["ccs_name"]])
    parent = None
    if "parent_code" in col and col["parent_code"] < len(row):
        p = _cell_str(row[col["parent_code"]])
        parent = (p[:64] if p else None) or None
    note = None
    if "ccs_note" in col and col["ccs_note"] < len(row):
        t = _cell_str(row[col["ccs_note"]])
        note = t or None
    return {
        "ccs_code": code,
        "ccs_name": name or None,
        "parent_code": parent,
        "ccs_note": note,
    }


def load_ics_rows_from_template_xlsx(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        first = next(rows_iter, None)
        if not first:
            return []
        col = _field_col_map(first, ICS_HEADER_TO_FIELD)
        if "ics_code" not in col:
            return []
        out: list[dict[str, Any]] = []
        for row in rows_iter:
            if not row:
                continue
            rec = _ics_row_from_values(row, col)
            if rec:
                out.append(rec)
        return out
    finally:
        wb.close()


def load_ccs_rows_from_template_xlsx(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        first = next(rows_iter, None)
        if not first:
            return []
        col = _field_col_map(first, CCS_HEADER_TO_FIELD)
        if "ccs_code" not in col:
            return []
        out: list[dict[str, Any]] = []
        for row in rows_iter:
            if not row:
                continue
            rec = _ccs_row_from_values(row, col)
            if rec:
                out.append(rec)
        return out
    finally:
        wb.close()


def load_ics_rows_from_template_csv(raw: bytes) -> list[dict[str, Any]]:
    text = raw.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    header = next(reader, None)
    if not header:
        return []
    col = _field_col_map(tuple(header), ICS_HEADER_TO_FIELD)
    if "ics_code" not in col:
        return []
    out: list[dict[str, Any]] = []
    for row in reader:
        if not row or not any(c.strip() for c in row if c):
            continue
        rec = _ics_row_from_values(tuple(row), col)
        if rec:
            out.append(rec)
    return out


def load_ccs_rows_from_template_csv(raw: bytes) -> list[dict[str, Any]]:
    text = raw.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    header = next(reader, None)
    if not header:
        return []
    col = _field_col_map(tuple(header), CCS_HEADER_TO_FIELD)
    if "ccs_code" not in col:
        return []
    out: list[dict[str, Any]] = []
    for row in reader:
        if not row or not any(c.strip() for c in row if c):
            continue
        rec = _ccs_row_from_values(tuple(row), col)
        if rec:
            out.append(rec)
    return out


def template_ics_csv_bytes() -> bytes:
    if ICS_TABLE_TEMPLATE_CSV.is_file():
        raw = ICS_TABLE_TEMPLATE_CSV.read_bytes()
        if raw.startswith(b"\xef\xbb\xbf"):
            return raw
        return b"\xef\xbb\xbf" + raw
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(list(ICS_TEMPLATE_HEADERS))
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def template_ccs_csv_bytes() -> bytes:
    if CCS_TABLE_TEMPLATE_CSV.is_file():
        raw = CCS_TABLE_TEMPLATE_CSV.read_bytes()
        if raw.startswith(b"\xef\xbb\xbf"):
            return raw
        return b"\xef\xbb\xbf" + raw
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(list(CCS_TEMPLATE_HEADERS))
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def template_ics_xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "ICS行业分类"
    for col, title in enumerate(ICS_TEMPLATE_HEADERS, start=1):
        ws.cell(row=1, column=col, value=title)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def template_ccs_xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "CCS行业分类"
    for col, title in enumerate(CCS_TEMPLATE_HEADERS, start=1):
        ws.cell(row=1, column=col, value=title)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def load_rows_from_template_upload(upload: UploadedFile, scheme: str) -> list[dict[str, Any]]:
    """从用户上传的 CSV / xlsx 解析为 ORM 可用的 dict 列表；scheme: ICS | CCS。"""
    sch = (scheme or "").strip().upper()
    if sch not in ("ICS", "CCS"):
        raise ValueError("scheme 必须是 ICS 或 CCS")
    name = (upload.name or "").lower()
    raw = upload.read()
    is_csv = name.endswith(".csv")

    if sch == "ICS":
        if is_csv:
            return load_ics_rows_from_template_csv(raw)
        from tempfile import NamedTemporaryFile

        with NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
            tmp.write(raw)
            p = Path(tmp.name)
        try:
            return load_ics_rows_from_template_xlsx(p)
        finally:
            p.unlink(missing_ok=True)

    if is_csv:
        return load_ccs_rows_from_template_csv(raw)
    from tempfile import NamedTemporaryFile

    with NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        tmp.write(raw)
        p = Path(tmp.name)
    try:
        return load_ccs_rows_from_template_xlsx(p)
    finally:
        p.unlink(missing_ok=True)
